from pathlib import Path
from io import BytesIO

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


def read_document(inbox, attachment_path):
    """
    Read TXT, PDF, DOCX and XLSX attachments
    and return their contents as plain text.
    """

    extension = Path(attachment_path).suffix.lower()

    # =========================================================
    # TXT
    # =========================================================

    if extension == ".txt":
        return inbox.read_text(attachment_path)

    # =========================================================
    # PDF
    # =========================================================

    if extension == ".pdf":
        data = inbox.read_bytes(attachment_path)
        reader = PdfReader(BytesIO(data))
        pages = []

        for page in reader.pages:
            text = page.extract_text()
            if text and text.strip():
                pages.append(text)

        result = "\n".join(pages).strip()

        if not result:
            raise ValueError("PDF contains no extractable text")

        return result

    # =========================================================
    # DOCX
    # =========================================================

    if extension == ".docx":

        data = inbox.read_bytes(attachment_path)

        document = Document(BytesIO(data))

        lines = []

        # Normal paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                lines.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:
                        values.append(value)

                if values:
                    lines.append(" | ".join(values))

        return "\n".join(lines)

    # =========================================================
    # XLSX
    # =========================================================

    if extension == ".xlsx":

        data = inbox.read_bytes(attachment_path)

        workbook = load_workbook(
            BytesIO(data),
            data_only=True
        )

        lines = []

        for sheet in workbook.worksheets:

            lines.append(f"--- SHEET: {sheet.title} ---")

            for row in sheet.iter_rows(values_only=True):

                values = []

                for value in row:

                    if value is not None:

                        values.append(str(value).strip())

                if values:

                    lines.append(" | ".join(values))

        return "\n".join(lines)

    # =========================================================
    # Unsupported
    # =========================================================

    raise ValueError(
        f"Unsupported document format: {extension}"
    )